from pathlib import Path

from docx import Document


INPUT = Path(r"D:\BOMP原型demo\outputs\BOMP知识共享中心业务需求说明书-4.4.2修订版.docx")
TEMP_OUTPUT = INPUT.with_name(INPUT.stem + ".producer.tmp.docx")


def set_cell_text(cell, value):
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for paragraph in cell.paragraphs[1:]:
        for run in paragraph.runs:
            run.text = ""


doc = Document(str(INPUT))
tables = [
    table
    for table in doc.tables
    if table.rows and any(cell.text.strip() == "制作单位" for cell in table.rows[0].cells)
]
if len(tables) != 1:
    raise RuntimeError(f"Expected one 4.6.1 business-information table, found {len(tables)}")

table = tables[0]
header = [cell.text.strip() for cell in table.rows[0].cells]
producer_column = header.index("制作单位")
for row in table.rows[1:]:
    set_cell_text(row.cells[producer_column], "供应链运营中心")

doc.save(str(TEMP_OUTPUT))
TEMP_OUTPUT.replace(INPUT)
print(f"updated_rows={len(table.rows) - 1}")
