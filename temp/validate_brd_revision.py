from pathlib import Path
import zipfile

from docx import Document


SOURCE = Path(r"D:\BOMP原型demo\outputs\BOMP知识共享中心业务需求说明书.docx")
OUTPUT = Path(r"D:\BOMP原型demo\outputs\BOMP知识共享中心业务需求说明书-4.4.2修订版.docx")


source_doc = Document(str(SOURCE))
output_doc = Document(str(OUTPUT))

activity_tables = [
    table
    for table in output_doc.tables
    if table.rows and table.cell(0, 0).text.strip() == "活动编号"
]
assert len(activity_tables) == 13
assert all(len(table.rows) == 3 for table in activity_tables)
assert all(
    table.cell(0, 0).text == "活动编号"
    and table.cell(0, 2).text == "活动名称"
    and table.cell(1, 2).text == "使用组织单元"
    and table.cell(1, 3).text == "国网物资公司"
    and table.cell(2, 0).text == "活动描述"
    for table in activity_tables
)

information_tables = [
    table
    for table in output_doc.tables
    if table.rows and any(cell.text.strip() == "制作单位" for cell in table.rows[0].cells)
]
assert len(information_tables) == 1
information_table = information_tables[0]
header = [cell.text.strip() for cell in information_table.rows[0].cells]
producer_column = header.index("制作单位")
assert len(information_table.rows) == 13
assert all(
    row.cells[producer_column].text == "供应链运营中心"
    for row in information_table.rows[1:]
)

assert [paragraph.text for paragraph in source_doc.paragraphs] == [
    paragraph.text for paragraph in output_doc.paragraphs
]
changed_non_target_tables = []
for index, (source_table, output_table) in enumerate(
    zip(source_doc.tables, output_doc.tables)
):
    if source_table.cell(0, 0).text.strip() == "流程编号" and len(source_table.rows) == 7:
        continue
    if any(cell.text.strip() == "制作单位" for cell in source_table.rows[0].cells):
        continue
    source_cells = [[cell.text for cell in row.cells] for row in source_table.rows]
    output_cells = [[cell.text for cell in row.cells] for row in output_table.rows]
    if source_cells != output_cells:
        changed_non_target_tables.append(index)
assert not changed_non_target_tables, changed_non_target_tables

with zipfile.ZipFile(OUTPUT) as archive:
    assert archive.testzip() is None
    settings = archive.read("word/settings.xml")
    document_xml = archive.read("word/document.xml")
    assert b"w:updateFields" in settings and b'w:val="true"' in settings
    assert b"TOC" in document_xml

print(
    "PASS 4.4.2_tables=13; 4.6.1_records=12; producer_values=12; "
    "non_target_content_unchanged=True; zip_integrity=True; "
    "update_fields=True; toc_field=True"
)
