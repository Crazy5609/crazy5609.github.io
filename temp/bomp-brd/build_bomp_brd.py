from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


REFERENCE = Path(
    r"C:\Users\admin\Desktop\BOMP\业务需求说明书\国网物资公司-2025年企业经营管理系统-设计开发实施项目-财务预算管理业务需求规格说明书.docx"
)
OUTPUT = Path(r"D:\BOMP原型demo\outputs\BOMP知识共享中心业务需求说明书.docx")


ROLES = {
    "R_BOMP_01": "员工",
    "R_BOMP_02": "部门知识库管理员",
    "R_BOMP_03": "部门知识库审核人",
    "R_BOMP_04": "专业管理部门知识库审核人",
    "R_BOMP_05": "法律事务部合规审查人",
    "R_BOMP_06": "综合管理部保密审查人",
    "R_BOMP_07": "公司知识中心管理员",
    "R_BOMP_08": "工作规则编制部门用户",
    "R_BOMP_09": "总法律顾问审核人",
}


PROCESS_ROWS = [
    {
        "id": "BP_BOMP_01",
        "name": "知识共享中心访问",
        "function": "共享中心首页各层级界面",
        "roles": "R_BOMP_01",
        "bs": "BS_BOMP_0101",
        "inputs": "P_BOMP_01、P_BOMP_12",
        "outputs": "P_BOMP_12",
        "description": "用户进入知识共享中心后，系统依据当前账号、岗位和角色加载首页栏目、重点内容、快捷入口及导航层级。用户展开导航或选择入口后，进入与所选层级一致且有权访问的页面。",
        "rules": "首页配置须处于生效状态；栏目、快捷入口和页面入口均须经过权限过滤；无权限入口和操作不予提供。",
    },
    {
        "id": "BP_BOMP_02",
        "name": "知识目录分类浏览",
        "function": "文件分级分类管理",
        "roles": "R_BOMP_01、R_BOMP_02、R_BOMP_07",
        "bs": "BS_BOMP_0201",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04",
        "outputs": "P_BOMP_02、P_BOMP_03",
        "description": "用户从公司级或部门级知识入口进入目录树，按目录层级逐级浏览。系统同步展示当前目录、完整目录层级以及该目录范围内有权访问的文件。",
        "rules": "一级目录汇总其全部下级目录文件，下级目录仅展示本目录范围；界面目录、目录记录和文件分区须保持一致。",
    },
    {
        "id": "BP_BOMP_03",
        "name": "公司知识中心管理",
        "function": "公司知识中心",
        "roles": "R_BOMP_01、R_BOMP_07",
        "bs": "BS_BOMP_0301",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04",
        "outputs": "P_BOMP_02",
        "description": "用户按公司知识目录浏览文件，查看文件详情，并在授权范围内执行下载、反馈、收藏和删除等操作；公司知识中心管理员可执行批量上传或批量下载。",
        "rules": "文件列表、详情和操作均须实时校验权限；删除须二次确认；批量操作按文件逐项校验并反馈处理结果。",
    },
    {
        "id": "BP_BOMP_04",
        "name": "部门知识库管理",
        "function": "部门知识库",
        "roles": "R_BOMP_01、R_BOMP_02",
        "bs": "BS_BOMP_0401",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04",
        "outputs": "P_BOMP_02",
        "description": "用户按部门知识目录浏览文件，查看文件详情并执行授权操作；部门知识库管理员可维护本部门目录和知识文件，并执行批量上传或批量下载。",
        "rules": "部门知识默认在本部门范围内展示，跨部门访问以实际授权为准；批量操作须逐项校验和反馈。",
    },
    {
        "id": "BP_BOMP_05",
        "name": "文件权限管理",
        "function": "文件权限管理",
        "roles": "R_BOMP_01、R_BOMP_02、R_BOMP_07",
        "bs": "BS_BOMP_0501",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_04",
        "outputs": "P_BOMP_04",
        "description": "管理人员查看文件当前授权范围，根据管理权限新增、调整或收回指定人员、部门的文件权限。用户访问列表、搜索结果、详情或下载文件时，系统重新校验权限。",
        "rules": "公司知识默认全员可见，部门知识默认本部门可见，个人知识默认仅创建人可见；转授权不得扩大操作者自身的管理范围。",
    },
    {
        "id": "BP_BOMP_06",
        "name": "知识审核与发布",
        "function": "知识审核管理",
        "roles": "R_BOMP_01、R_BOMP_02、R_BOMP_03、R_BOMP_04、R_BOMP_05、R_BOMP_06、R_BOMP_07",
        "bs": "BS_BOMP_0601",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_05",
        "outputs": "P_BOMP_02、P_BOMP_05",
        "description": "提交人完成文件上传、必要信息填写和共享目标选择后，系统按目标知识库生成审核任务。各环节审核人依次处理，全部通过后发布；退回后由提交人修改并重新提交。",
        "rules": "待审核文件不得向普通用户发布；任务只能由当前环节有权审核人处理；退回原因、修改内容和各轮审核记录均须保留。",
    },
    {
        "id": "BP_BOMP_07",
        "name": "知识检索",
        "function": "知识检索",
        "roles": "R_BOMP_01",
        "bs": "BS_BOMP_0701 至 BS_BOMP_0705",
        "inputs": "P_BOMP_01、P_BOMP_04、P_BOMP_06",
        "outputs": "P_BOMP_06",
        "description": "用户输入关键词或选择上传日期、上传部门、文件归属部门、文件类型、公开范围和知识类型等条件。系统在权限范围内执行关键词、正文、模糊或组合条件检索并返回结果。",
        "rules": "多个结构化条件采用同时满足关系；检索过程和结果展示均须执行权限过滤；无权文件不得泄露正文和附件内容。",
    },
    {
        "id": "BP_BOMP_08",
        "name": "临时共享网盘使用",
        "function": "临时共享网盘",
        "roles": "R_BOMP_01",
        "bs": "BS_BOMP_0801",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04",
        "outputs": "P_BOMP_02、P_BOMP_03",
        "description": "用户进入临时共享网盘，按目录层级浏览个人文件和群组文件，并在授权范围内执行上传、下载、移动、删除或创建文件夹等操作。",
        "rules": "个人文件中心按个人权限操作；群组文件中心按成员和文件夹权限操作；父文件夹管理员继承子文件夹管理权限。",
    },
    {
        "id": "BP_BOMP_09",
        "name": "多层级文件夹申请",
        "function": "多层级文件夹管理",
        "roles": "R_BOMP_01、R_BOMP_02、R_BOMP_07",
        "bs": "BS_BOMP_0901",
        "inputs": "P_BOMP_01、P_BOMP_03、P_BOMP_04、P_BOMP_07",
        "outputs": "P_BOMP_03、P_BOMP_07",
        "description": "申请人选择公司级或部门级目标路径，填写文件夹名称和创建原因并提交申请。系统校验后生成审批任务，审批通过后在目标路径创建文件夹。",
        "rules": "目标路径、必填项和同级名称须通过校验；审批流程依据目标目录层级确定；退回申请允许修改后重新提交。",
    },
    {
        "id": "BP_BOMP_10",
        "name": "个人上传与审批管理",
        "function": "个人中心",
        "roles": "R_BOMP_01、R_BOMP_03、R_BOMP_04、R_BOMP_05、R_BOMP_06",
        "bs": "BS_BOMP_1001",
        "inputs": "P_BOMP_01、P_BOMP_02、P_BOMP_05、P_BOMP_08",
        "outputs": "P_BOMP_08",
        "description": "用户通过个人中心查询本人上传文件及处理状态；审核人查看待处理和已处理任务，处理当前环节任务并查询历史记录。",
        "rules": "我的上传仅归集当前账号上传的文件；已处理任务不得重复操作；页面操作随文件和审核状态变化。",
    },
    {
        "id": "BP_BOMP_11",
        "name": "工作规则备案",
        "function": "工作规则备案",
        "roles": "R_BOMP_05、R_BOMP_06、R_BOMP_08、R_BOMP_09",
        "bs": "BS_BOMP_1101",
        "inputs": "P_BOMP_01、P_BOMP_09",
        "outputs": "P_BOMP_09",
        "description": "编制部门用户发起工作规则新增、更新或废止备案，申请依次流转至综合管理部、法律事务部和总法律顾问审核。全部通过后更新规则状态。",
        "rules": "审核须按既定顺序流转；退回时保留原规则状态和退回意见，申请人修改后可重新提交。",
    },
    {
        "id": "BP_BOMP_12",
        "name": "重复文件校验",
        "function": "重复文件校验",
        "roles": "R_BOMP_01、R_BOMP_02、R_BOMP_07",
        "bs": "BS_BOMP_1201",
        "inputs": "P_BOMP_01、P_BOMP_02",
        "outputs": "P_BOMP_10",
        "description": "用户选择待上传文件后，系统提取文件名称、摘要等可比对信息，与已有知识文件进行比对，并返回完全相同、高度相似或未发现重复的结果。",
        "rules": "命中重复或相似文件时须展示必要判断信息；判定标准由系统配置；未命中时继续正常上传流程。",
    },
    {
        "id": "BP_BOMP_13",
        "name": "群组文件管理",
        "function": "群组文件管理",
        "roles": "R_BOMP_01",
        "bs": "BS_BOMP_1301",
        "inputs": "P_BOMP_01、P_BOMP_04、P_BOMP_11",
        "outputs": "P_BOMP_04、P_BOMP_11",
        "description": "群组文件夹管理员维护成员并分配查看、上传、下载和删除等权限。普通成员在授权文件夹范围内操作文件，父文件夹管理员自动继承子文件夹管理权限。",
        "rules": "普通成员不得操作未授权内容；成员可访问授权文件夹向上至群组的路径；继承关系随层级或管理员身份变化重新计算。",
    },
]


ACTIVITY_ROWS = [
    ("BS_BOMP_0101", "共享中心首页访问", "BP_BOMP_01", "R_BOMP_01", "无", "加载首页栏目、快捷入口和导航层级。", "用户已登录且首页配置已生效。", "P_BOMP_01、P_BOMP_12", "P_BOMP_12", "F_BOMP_0101", "UC_BOMP_0101、UC_BOMP_0102"),
    ("BS_BOMP_0201", "文件分级分类浏览", "BP_BOMP_02", "R_BOMP_01、R_BOMP_02、R_BOMP_07", "BS_BOMP_0301、BS_BOMP_0401", "按公司级、部门级和下级目录展示对应文件。", "目录、文件分区和权限信息已配置。", "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04", "P_BOMP_02、P_BOMP_03", "F_BOMP_0201", "UC_BOMP_0201、UC_BOMP_0202"),
    ("BS_BOMP_0301", "公司知识中心管理", "BP_BOMP_03", "R_BOMP_01、R_BOMP_07", "无", "浏览公司知识并执行查看、下载、反馈、收藏、删除及批量操作。", "公司目录和文件已发布，用户具有相应权限。", "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04", "P_BOMP_02", "F_BOMP_0301", "UC_BOMP_0301、UC_BOMP_0302、UC_BOMP_0303、UC_BOMP_0304"),
    ("BS_BOMP_0401", "部门知识库管理", "BP_BOMP_04", "R_BOMP_01、R_BOMP_02", "无", "浏览部门知识并执行查看、下载、反馈、收藏、删除及批量操作。", "部门目录和文件已发布，用户具有相应权限。", "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04", "P_BOMP_02", "F_BOMP_0401", "UC_BOMP_0401、UC_BOMP_0402、UC_BOMP_0403"),
    ("BS_BOMP_0501", "文件权限管理", "BP_BOMP_05", "R_BOMP_01、R_BOMP_02、R_BOMP_07", "BS_BOMP_0301、BS_BOMP_0401", "查看、增加、调整或收回文件权限，并控制越权访问。", "操作者具有目标文件的权限管理权。", "P_BOMP_01、P_BOMP_02、P_BOMP_04", "P_BOMP_04", "F_BOMP_0501", "UC_BOMP_0501、UC_BOMP_0502、UC_BOMP_0503、UC_BOMP_0504"),
    ("BS_BOMP_0601", "知识审核管理", "BP_BOMP_06", "R_BOMP_01、R_BOMP_02、R_BOMP_03、R_BOMP_04、R_BOMP_05、R_BOMP_06、R_BOMP_07", "BS_BOMP_1201", "按共享目标执行部门、专业、合规和保密审核，处理退回重提并发布知识。", "文件已上传并完成必要信息填写。", "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_05", "P_BOMP_02、P_BOMP_05", "F_BOMP_0601", "UC_BOMP_0601、UC_BOMP_0602、UC_BOMP_0603、UC_BOMP_0604、UC_BOMP_0605、UC_BOMP_0606"),
    ("BS_BOMP_0701", "关键词检索", "BP_BOMP_07", "R_BOMP_01", "无", "按关键词检索文件名称、基础信息和可检索正文。", "用户已登录，检索索引已建立。", "P_BOMP_01、P_BOMP_04、P_BOMP_06", "P_BOMP_06", "F_BOMP_0701", "UC_BOMP_0701"),
    ("BS_BOMP_0702", "正文穿透检索", "BP_BOMP_07", "R_BOMP_01", "无", "在权限范围内检索知识文件正文并返回命中位置。", "文件正文已完成解析和索引。", "P_BOMP_01、P_BOMP_04、P_BOMP_06", "P_BOMP_06", "F_BOMP_0701", "UC_BOMP_0701"),
    ("BS_BOMP_0703", "模糊匹配检索", "BP_BOMP_07", "R_BOMP_01", "无", "对关键词及相近文本进行模糊匹配并返回相关文件。", "用户已输入检索词。", "P_BOMP_01、P_BOMP_04、P_BOMP_06", "P_BOMP_06", "F_BOMP_0701", "UC_BOMP_0701"),
    ("BS_BOMP_0704", "高级组合条件查询", "BP_BOMP_07", "R_BOMP_01", "无", "按上传日期、部门、类型、公开范围和知识类型等条件组合查询。", "用户至少填写一个有效查询条件。", "P_BOMP_01、P_BOMP_04、P_BOMP_06", "P_BOMP_06", "F_BOMP_0701", "UC_BOMP_0702"),
    ("BS_BOMP_0705", "检索结果权限过滤", "BP_BOMP_07", "R_BOMP_01", "BS_BOMP_0701、BS_BOMP_0702、BS_BOMP_0703、BS_BOMP_0704", "按当前用户权限过滤检索结果并控制结果摘要和文件访问。", "检索已返回候选结果。", "P_BOMP_01、P_BOMP_04、P_BOMP_06", "P_BOMP_06", "F_BOMP_0701", "UC_BOMP_0701、UC_BOMP_0702"),
    ("BS_BOMP_0801", "临时共享网盘使用", "BP_BOMP_08", "R_BOMP_01", "无", "按权限浏览和操作个人文件、群组文件及多层级目录。", "用户具有临时共享网盘访问权限。", "P_BOMP_01、P_BOMP_02、P_BOMP_03、P_BOMP_04", "P_BOMP_02、P_BOMP_03", "F_BOMP_0801", "UC_BOMP_0801"),
    ("BS_BOMP_0901", "多层级文件夹申请", "BP_BOMP_09", "R_BOMP_01、R_BOMP_02、R_BOMP_07", "BS_BOMP_0801", "提交新建文件夹申请，审批通过后在目标路径创建目录。", "申请人具有发起权限且目标路径有效。", "P_BOMP_01、P_BOMP_03、P_BOMP_04、P_BOMP_07", "P_BOMP_03、P_BOMP_07", "F_BOMP_0901", "UC_BOMP_0901"),
    ("BS_BOMP_1001", "个人上传与审批管理", "BP_BOMP_10", "R_BOMP_01、R_BOMP_03、R_BOMP_04、R_BOMP_05、R_BOMP_06", "BS_BOMP_0501、BS_BOMP_0601", "查询个人上传文件和审核任务，处理当前环节任务并查看历史。", "当前账号存在上传文件或审核任务。", "P_BOMP_01、P_BOMP_02、P_BOMP_05、P_BOMP_08", "P_BOMP_08", "F_BOMP_1001", "UC_BOMP_1001、UC_BOMP_1002"),
    ("BS_BOMP_1101", "工作规则备案", "BP_BOMP_11", "R_BOMP_05、R_BOMP_06、R_BOMP_08、R_BOMP_09", "无", "发起工作规则新增、更新或废止备案并完成分级审核。", "编制部门用户具有备案权限。", "P_BOMP_01、P_BOMP_09", "P_BOMP_09", "F_BOMP_1101", "UC_BOMP_1101"),
    ("BS_BOMP_1201", "重复文件校验", "BP_BOMP_12", "R_BOMP_01、R_BOMP_02、R_BOMP_07", "无", "比对待上传文件与已有文件并提示完全相同或高度相似结果。", "用户具有上传权限并已选择文件。", "P_BOMP_01、P_BOMP_02", "P_BOMP_10", "F_BOMP_1201", "UC_BOMP_1201"),
    ("BS_BOMP_1301", "群组文件管理", "BP_BOMP_13", "R_BOMP_01", "无", "维护群组文件夹成员、权限和父子文件夹权限继承。", "群组和目标文件夹已创建。", "P_BOMP_01、P_BOMP_04、P_BOMP_11", "P_BOMP_04、P_BOMP_11", "F_BOMP_1301", "UC_BOMP_1301、UC_BOMP_1302"),
]


INFO_ROWS = [
    ("P_BOMP_01", "表单", "用户、岗位与角色信息", "身份识别及权限匹配", "所有部门（单位）", "所有部门（单位）", "实时", "用户标识、姓名、所属组织、岗位、角色、账号状态", "用户身份、组织和角色共同决定功能及数据访问范围。"),
    ("P_BOMP_02", "表单", "知识文件基础信息", "展示和管理文件名称、类型、上传人、时间、归属部门等", "所有部门（单位）", "文件上传人", "按次", "文件标识、名称、类型、上传人、上传时间、归属部门、知识类型、版本、状态", "文件状态与目录、权限和审核状态保持一致。"),
    ("P_BOMP_03", "表单", "知识目录信息", "管理公司级、部门级及文件夹层级", "所有部门（单位）", "知识库管理员", "按需", "目录标识、目录名称、父目录、完整路径、目录层级、知识库类型", "目录层级与文件分区一一对应，同级目录名称不得重复。"),
    ("P_BOMP_04", "表单", "文件权限信息", "控制审批、上传、查看、下载、删除、收藏等操作", "所有部门（单位）", "知识库管理员", "按需", "文件标识、授权对象、对象类型、权限项、权限来源、有效状态", "权限不得超出授权人管理范围，访问入口统一校验。"),
    ("P_BOMP_05", "表单", "知识审核任务与记录", "记录提交、审核、退回、发布及各环节状态", "相关审核部门", "系统", "按次", "申请标识、文件、目标知识库、当前环节、状态、处理人、意见、处理时间、审核轮次", "审核按配置顺序流转，历史记录不得覆盖。"),
    ("P_BOMP_06", "表单", "知识检索条件与结果", "执行正文、模糊及组合条件查询", "所有部门（单位）", "用户或系统", "高频", "关键词、检索范围、日期范围、上传部门、归属部门、文件类型、公开范围、知识类型、结果数量", "组合条件同时满足，检索和结果展示均执行权限过滤。"),
    ("P_BOMP_07", "表单", "文件夹申请信息", "记录目标路径、名称、创建原因和审批状态", "所有部门（单位）", "申请人", "按次", "申请标识、目标路径、文件夹名称、创建原因、申请人、审批状态、审批意见", "目标路径有效且同级名称唯一，审批通过后创建目录。"),
    ("P_BOMP_08", "表单", "个人上传与审批信息", "汇总个人文件状态和个人审批任务", "所有部门（单位）", "用户或系统", "实时", "用户、上传文件、处理状态、审核进度、退回原因、待办任务、已办任务", "仅展示当前账号上传记录或当前角色有权处理的任务。"),
    ("P_BOMP_09", "表单", "工作规则备案信息", "管理新增、更新、废止及审批记录", "编制部门及审核部门", "编制部门", "按次", "备案类型、规则名称、规则文件、申请说明、申请人、审核状态、审核意见、规则状态", "按综合管理、法律审查和总法律顾问终审顺序流转。"),
    ("P_BOMP_10", "表单", "重复文件比对结果", "提示高度相似或完全相同文件", "所有部门（单位）", "系统", "上传时", "待上传文件、命中文件、判定类型、判断依据、比对时间", "判定标准由系统配置，未命中时允许继续上传。"),
    ("P_BOMP_11", "表单", "群组成员与文件夹权限信息", "管理人员、角色、继承权限和路径访问", "相关群组", "群组管理员", "按需", "群组、文件夹、成员、成员角色、权限项、权限来源、继承关系、路径范围", "父文件夹管理员继承子文件夹管理权限，普通成员仅操作授权内容。"),
    ("P_BOMP_12", "配置信息", "知识共享中心首页配置信息", "配置首页栏目、重点内容和快捷入口", "所有部门（单位）", "系统", "按需", "栏目名称、内容类型、显示顺序、入口目标、生效状态、适用范围", "仅加载已生效且当前用户有权访问的栏目和入口。"),
]


def set_run_font(run, name="宋体", size=Pt(10.5), bold=None):
    run.font.name = name
    run.font.size = size
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), "Times New Roman" if name == "宋体" else name)
    rfonts.set(qn("w:hAnsi"), "Times New Roman" if name == "宋体" else name)
    rfonts.set(qn("w:eastAsia"), name)


def clear_container(container):
    element = container._element
    for child in list(element):
        element.remove(child)
    p = OxmlElement("w:p")
    element.append(p)


def clear_body(doc):
    body = doc._element.body
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def ensure_heading_style(doc, level, size, before, after):
    name = f"BOMP Heading {level}"
    if name in doc.styles:
        style = doc.styles[name]
    else:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "黑体"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True
    pf.line_spacing = 1.3
    ppr = style.element.get_or_add_pPr()
    outline = ppr.find(qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        ppr.append(outline)
    outline.set(qn("w:val"), str(level - 1))
    return style


def configure_document(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)
    ensure_heading_style(doc, 1, 14, 12, 8)
    ensure_heading_style(doc, 2, 12, 10, 6)
    ensure_heading_style(doc, 3, 12, 8, 5)
    ensure_heading_style(doc, 4, 10.5, 6, 4)
    for level in range(1, 5):
        style_name = f"toc {level}"
        if style_name in doc.styles:
            toc_style = doc.styles[style_name]
        else:
            toc_style = doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        toc_style.font.name = "宋体"
        toc_style.font.size = Pt(9.5)
        toc_style.font.color.rgb = RGBColor(0, 0, 0)
        toc_style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        toc_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        toc_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        toc_style.paragraph_format.line_spacing = 1.0
        toc_style.paragraph_format.space_before = Pt(0)
        toc_style.paragraph_format.space_after = Pt(0)
    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = Mm(31.75)
        section.right_margin = Mm(31.75)
        section.top_margin = Mm(25.4)
        section.bottom_margin = Mm(25.4)
        section.header_distance = Mm(15)
        section.footer_distance = Mm(17.5)
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        clear_container(section.header)
        clear_container(section.footer)
        p = section.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run("第")
        set_run_font(run, "宋体", Pt(9))
        add_field(p, "PAGE")
        run = p.add_run("页")
        set_run_font(run, "宋体", Pt(9))


def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    return run


def add_heading(doc, number, title, level=1):
    p = doc.add_paragraph(style=f"BOMP Heading {level}")
    p.add_run(f"{number} {title}")
    return p


def add_body(doc, text="", indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    if text:
        set_run_font(p.add_run(text))
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margin(cell, top=70, start=80, bottom=70, end=80):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="B7B7B7", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, size=9, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(str(text))
    set_run_font(run, "宋体", Pt(size), bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)
    if fill:
        shade_cell(cell, fill)


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(7.5)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    set_run_font(p.add_run(text), "黑体", Pt(9))
    return p


def add_grid_table(doc, headers, rows, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, value in enumerate(headers):
        header.cells[i].width = Mm(widths[i])
        set_cell_text(header.cells[i], value, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=font_size, fill="BFBFBF")
    for values in rows:
        row = table.add_row()
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for i, value in enumerate(values):
            row.cells[i].width = Mm(widths[i])
            align = WD_ALIGN_PARAGRAPH.CENTER if i < 4 and len(str(value)) < 45 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(row.cells[i], value, align=align, size=font_size)
    return table


def add_process_detail(doc, process, index):
    add_heading(doc, f"4.4.2.{index}", process["name"], 4)
    table = doc.add_table(rows=7, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [26, 47, 26, 47]
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Mm(widths[i])
    pairs = [
        ("流程编号", process["id"], "流程名称", process["name"]),
        ("参与角色", process["roles"], "对应业务活动", process["bs"]),
    ]
    for r, values in enumerate(pairs):
        for c, value in enumerate(values):
            is_label = c in (0, 2)
            set_cell_text(table.cell(r, c), value, bold=is_label, align=WD_ALIGN_PARAGRAPH.CENTER if is_label else WD_ALIGN_PARAGRAPH.LEFT, fill="D9D9D9" if is_label else None)
    labels = [
        (2, "流程描述", process["description"]),
        (3, "输入业务信息编号", process["inputs"]),
        (4, "输出业务信息编号", process["outputs"]),
        (5, "业务规则", process["rules"]),
        (6, "对应业务职能", process["function"]),
    ]
    for row_index, label, value in labels:
        set_cell_text(table.cell(row_index, 0), label, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9D9D9")
        merged = table.cell(row_index, 1).merge(table.cell(row_index, 3))
        set_cell_text(merged, value, align=WD_ALIGN_PARAGRAPH.LEFT)
    return table


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    set_run_font(p.add_run("国网物资有限公司企业经营管理系统"), "黑体", Pt(24), True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    set_run_font(p.add_run("BOMP知识共享中心业务需求说明书"), "黑体", Pt(26), True)
    for _ in range(7):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("2026年9月"), "宋体", Pt(16), True)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_front_matter(doc):
    table = doc.add_table(rows=3, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    control = [
        ("文件编号", "", "生效日期", "", "受控编号", ""),
        ("密级", "", "版次", "V1.0", "修改状态", "初稿"),
        ("总页数", "", "正文", "", "附录", ""),
    ]
    widths = [23, 26, 23, 26, 23, 25]
    for r, values in enumerate(control):
        for c, value in enumerate(values):
            table.cell(r, c).width = Mm(widths[c])
            label = c % 2 == 0
            set_cell_text(table.cell(r, c), value, bold=label, align=WD_ALIGN_PARAGRAPH.CENTER, fill="D9D9D9" if label else None)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("文档修改记录"), "黑体", Pt(16), True)
    add_grid_table(doc, ["版本", "日期", "修改页", "作者", "批准人"], [("V1.0", "2026年9月", "全部", "", "")], [22, 30, 28, 33, 33], 9)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    set_run_font(p.add_run("目录"), "黑体", Pt(16), True)
    toc = doc.add_paragraph()
    add_field(toc, 'TOC \\o "1-4" \\h \\z \\u', "打开文档后更新目录")
    toc.add_run().add_break(WD_BREAK.PAGE)


def add_chapter_four(doc):
    add_heading(doc, "4", "业务描述", 1)
    add_heading(doc, "4.1", "业务目标", 2)
    add_body(doc)
    add_heading(doc, "4.2", "管理模式", 2)
    add_body(doc)
    add_heading(doc, "4.3", "组织单元", 2)
    add_body(doc)
    add_heading(doc, "4.4", "业务流程", 2)
    add_heading(doc, "4.4.1", "流程清单", 3)
    add_caption(doc, "表 4-1 BOMP知识共享中心业务流程清单")
    add_grid_table(
        doc,
        ["流程编号", "业务流程名称", "父级流程编号", "业务职能名称"],
        [(p["id"], p["name"], "", p["function"]) for p in PROCESS_ROWS],
        [30, 42, 30, 44],
        8.5,
    )
    add_heading(doc, "4.4.2", "业务流程分项说明", 3)
    for i, process in enumerate(PROCESS_ROWS, start=1):
        add_process_detail(doc, process, i)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "4.5", "业务活动", 2)
    add_heading(doc, "4.5.1", "业务活动清单", 3)
    add_caption(doc, "表 4-2 BOMP知识共享中心业务活动清单")
    add_grid_table(
        doc,
        ["业务活动编号", "业务活动名称", "使用角色编号", "依赖业务活动编号", "业务活动内容描述", "前置条件"],
        [(r[0], r[1], r[3], r[4], r[5], r[6]) for r in ACTIVITY_ROWS],
        [25, 25, 23, 25, 30, 18],
        7.5,
    )
    add_heading(doc, "4.5.2", "业务活动分项说明", 3)
    add_caption(doc, "表 4-3 BOMP知识共享中心业务步骤清单")
    add_grid_table(
        doc,
        ["业务步骤编号", "业务步骤名称", "输入业务信息编号", "输出业务信息编号", "业务步骤内容描述", "前置条件"],
        [(r[0], r[1], r[7], r[8], r[5], r[6]) for r in ACTIVITY_ROWS],
        [25, 25, 24, 22, 32, 18],
        7.5,
    )
    add_heading(doc, "4.6", "业务信息", 2)
    add_heading(doc, "4.6.1", "BOMP知识共享中心", 3)
    add_caption(doc, "表 4-4 BOMP知识共享中心业务信息清单")
    add_grid_table(
        doc,
        ["业务信息编号", "类型", "业务信息名称", "用途", "使用单位", "制作单位", "使用频率"],
        [row[:7] for row in INFO_ROWS],
        [22, 14, 25, 31, 20, 20, 14],
        7.5,
    )


def add_chapter_five(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "5", "共享融合需求分析", 1)
    add_body(doc)


def add_chapter_six(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, "6", "附录", 1)
    add_heading(doc, "6.1", "业务信息详单", 2)
    add_heading(doc, "6.1.1", "BOMP知识共享中心", 3)
    add_caption(doc, "表 6-1 BOMP知识共享中心业务信息详单")
    add_grid_table(
        doc,
        ["业务信息编号", "业务信息名称", "关键业务信息项", "主要业务规则"],
        [(row[0], row[2], row[7], row[8]) for row in INFO_ROWS],
        [24, 32, 47, 43],
        8,
    )
    add_heading(doc, "6.2", "需求跟踪矩阵", 2)
    add_caption(doc, "表 6-2 BOMP知识共享中心需求跟踪矩阵")
    add_grid_table(
        doc,
        ["业务流程编号", "业务活动编号", "软件功能点编号", "用例编号", "输入业务信息编号", "输出业务信息编号"],
        [(row[2], row[0], row[9], row[10], row[7], row[8]) for row in ACTIVITY_ROWS],
        [23, 24, 23, 28, 25, 23],
        7.5,
    )


def enable_field_updates(doc):
    settings = doc.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_body(doc)
    configure_document(doc)
    doc.core_properties.title = "BOMP知识共享中心业务需求说明书"
    doc.core_properties.subject = "BOMP知识共享中心业务需求"
    doc.core_properties.comments = ""
    add_cover(doc)
    add_front_matter(doc)
    add_chapter_four(doc)
    add_chapter_five(doc)
    add_chapter_six(doc)
    enable_field_updates(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
