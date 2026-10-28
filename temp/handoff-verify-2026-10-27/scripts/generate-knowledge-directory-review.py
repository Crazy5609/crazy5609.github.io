from __future__ import annotations

from collections import OrderedDict
import argparse
import json
import re
from pathlib import Path

import openpyxl


PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = PROJECT_ROOT / "知识目录"
OUTPUT_DIR = PROJECT_ROOT / "outputs" / "review" / "knowledge-directories"


def clean(value: object) -> str:
    if value is None:
        return ""
    text = " ".join(str(value).replace("\r", "\n").split())
    return "" if text == "/" else text


def merged_value_map(sheet) -> dict[tuple[int, int], object]:
    values: dict[tuple[int, int], object] = {}
    for merged_range in sheet.merged_cells.ranges:
        value = sheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                values[(row, column)] = value
    return values


def extract_paths(sheet, first_column: int, last_column: int) -> list[tuple[str, ...]]:
    merged_values = merged_value_map(sheet)
    paths: list[tuple[str, ...]] = []
    seen: set[tuple[str, ...]] = set()
    for row in range(3, sheet.max_row + 1):
        levels = []
        for column in range(first_column, last_column + 1):
            value = sheet.cell(row, column).value
            if value is None:
                value = merged_values.get((row, column))
            normalized = clean(value)
            if normalized:
                levels.append(normalized)
        path = tuple(levels)
        if path and path not in seen:
            paths.append(path)
            seen.add(path)
    return paths


def extract_paths_with_counts(
    sheet, first_column: int, last_column: int, knowledge_column: int
) -> tuple[list[tuple[str, ...]], dict[tuple[str, ...], int], dict[tuple[str, ...], int]]:
    merged_values = merged_value_map(sheet)
    paths: list[tuple[str, ...]] = []
    path_rows: dict[tuple[str, ...], int] = {}
    counts: dict[tuple[str, ...], int] = {}
    for row in range(3, sheet.max_row + 1):
        levels = []
        for column in range(first_column, last_column + 1):
            value = sheet.cell(row, column).value
            if value is None:
                value = merged_values.get((row, column))
            normalized = clean(value)
            if normalized:
                levels.append(normalized)
        path = tuple(levels)
        if not path:
            continue
        if path not in path_rows:
            paths.append(path)
            path_rows[path] = row
        knowledge = clean(sheet.cell(row, knowledge_column).value)
        if knowledge:
            for depth in range(1, len(path) + 1):
                ancestor = path[:depth]
                counts[ancestor] = counts.get(ancestor, 0) + 1
    return paths, path_rows, counts


DEPARTMENT_KEY_OVERRIDES = {
    ("架构管理",): "dept-architecture-management",
    ("架构管理", "供应链业务架构管理"): "dept-supply-chain-architecture",
    (
        "数字化项目建设管理",
        "综合计划项目",
        "自建系统（通过接口与IRS集成，不涉及数据重复录入问题）",
        "国网物资公司-2025年企业经营管理系统-设计开发实施项目",
    ): "dept-bomp-2025-project",
}


def build_sidebar_tree(
    paths: list[tuple[str, ...]],
    path_rows: dict[tuple[str, ...], int],
    counts: dict[tuple[str, ...], int],
) -> list[dict]:
    tree = build_tree(paths)

    def convert(branch: OrderedDict, parent: tuple[str, ...] = ()) -> list[dict]:
        nodes = []
        for label, children in branch.items():
            path = parent + (label,)
            row = path_rows.get(path, path_rows.get(next((item for item in paths if item[: len(path)] == path), path), 0))
            node = {
                "key": DEPARTMENT_KEY_OVERRIDES.get(path, f"dept-directory-r{row}-l{len(path)}"),
                "label": label,
                "count": counts.get(path, 0),
            }
            child_nodes = convert(children, path)
            if child_nodes:
                node["children"] = child_nodes
            nodes.append(node)
        return nodes

    return convert(tree)


def render_sidebar_data(nodes: list[dict]) -> str:
    payload = json.dumps(nodes, ensure_ascii=False, indent=2)
    indented = "\n".join(f"  {line}" for line in payload.splitlines())
    return (
        "  // DEPARTMENT_DIRECTORY_START - generated from the approved department directory workbook.\n"
        f"  const departmentDirectories = {indented.lstrip()};\n\n"
        "  window.BompDepartmentDirectoryData = departmentDirectories;\n"
        "  // DEPARTMENT_DIRECTORY_END"
    )


def update_sidebar_component(nodes: list[dict]) -> None:
    component_file = PROJECT_ROOT / "components" / "app-sidebar" / "component.js"
    source = component_file.read_text(encoding="utf-8")
    replacement = render_sidebar_data(nodes)
    updated, replacements = re.subn(
        r"  // DEPARTMENT_DIRECTORY_START[\s\S]*?  // DEPARTMENT_DIRECTORY_END",
        replacement,
        source,
        count=1,
    )
    if replacements != 1:
        raise RuntimeError("Department directory markers were not found exactly once.")
    component_file.write_text(updated, encoding="utf-8", newline="\n")


def build_tree(paths: list[tuple[str, ...]]) -> OrderedDict:
    tree: OrderedDict = OrderedDict()
    for path in paths:
        branch = tree
        for label in path:
            branch = branch.setdefault(label, OrderedDict())
    return tree


def count_levels(tree: OrderedDict, depth: int = 0, counts: list[int] | None = None) -> list[int]:
    if counts is None:
        counts = []
    for children in tree.values():
        if len(counts) <= depth:
            counts.append(0)
        counts[depth] += 1
        count_levels(children, depth + 1, counts)
    return counts


def render_tree(tree: OrderedDict, depth: int = 0) -> list[str]:
    lines: list[str] = []
    for label, children in tree.items():
        lines.append(f"{'  ' * depth}- {label}")
        lines.extend(render_tree(children, depth + 1))
    return lines


def render_document(
    title: str,
    source_name: str,
    sheet_name: str,
    source_columns: str,
    ignored_column: str,
    paths: list[tuple[str, ...]],
    root_label: str | None = None,
    notes: list[str] | None = None,
) -> str:
    tree = build_tree(paths)
    counts = count_levels(tree)
    displayed_tree = OrderedDict([(root_label, tree)]) if root_label else tree
    lines = [
        f"# {title}",
        "",
        "> 状态：已确认并应用。当前目录口径已同步到项目公共侧栏及共用目录数据。",
        "",
        "## 来源与提取规则",
        "",
        f"- 来源文件：`知识目录/{source_name}`",
        f"- 来源工作表：`{sheet_name}`",
        f"- 目录层级列：`{source_columns}`",
        f"- 不生成节点：空值、`/`、{ignored_column}",
        "- 同一目录路径按首次出现顺序保留，合并单元格按其左上角值向下还原。",
        "",
        "## 目录统计",
        "",
        f"- 来源目录路径：{len(paths)} 条",
        f"- 一级目录：{counts[0] if counts else 0} 个",
        f"- 各层节点数：{' / '.join(f'{index + 1}级 {count}' for index, count in enumerate(counts))}",
        f"- 目录节点总数：{sum(counts)} 个" + (f"（不含根节点“{root_label}”）" if root_label else ""),
    ]
    if notes:
        lines.extend(["", "## 口径说明", ""] + [f"- {note}" for note in notes])
    lines.extend(["", "## 完整目录", "", *render_tree(displayed_tree), ""])
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--print-department-sidebar", action="store_true")
    parser.add_argument("--update-sidebar", action="store_true")
    args = parser.parse_args()
    company_file = SOURCE_DIR / "公司级.xlsx"
    department_file = next(SOURCE_DIR.glob("部门级*.xlsx"))

    company_sheet = openpyxl.load_workbook(company_file, data_only=True, read_only=False)["公司知识库目录层级"]
    department_sheet = openpyxl.load_workbook(department_file, data_only=True, read_only=False)["部门知识库样例"]

    company_paths = extract_paths(company_sheet, 2, 7)
    department_paths, department_path_rows, department_counts = extract_paths_with_counts(
        department_sheet, 2, 5, 6
    )

    department_sidebar = build_sidebar_tree(
        department_paths, department_path_rows, department_counts
    )

    if args.print_department_sidebar:
        print(json.dumps(
            department_sidebar,
            ensure_ascii=False,
            indent=2,
        ))
        return

    if args.update_sidebar:
        update_sidebar_component(department_sidebar)
        return

    company_markdown = render_document(
        title="公司级知识目录（评审稿）",
        source_name=company_file.name,
        sheet_name=company_sheet.title,
        source_columns="B-G（一级目录、二级目录、1级文件夹、2级文件夹、3级文件夹、4…N级文件夹）",
        ignored_column="H 列“知识资料内容”",
        paths=company_paths,
        notes=[
            "该文件是完整公司级目录的主数据源。",
            "G 列当前均为空值或 `/`，因此实际目录最深为五级，但仍保留继续扩展 4…N 级文件夹的口径。",
        ],
    )
    department_markdown = render_document(
        title="部门级知识目录（供应链运营中心，评审稿）",
        source_name=department_file.name,
        sheet_name=department_sheet.title,
        source_columns="B-E（一级目录、二级目录、一级文件夹、二级文件夹）",
        ignored_column="F 列“知识文件”",
        paths=department_paths,
        root_label="供应链运营中心",
        notes=[
            "G 列为空白分隔列，H-K 列用于表达部门知识共享到公司级知识库后的目标映射，不作为部门目录节点。",
            "当前文件实际包含 14 个部门一级目录；后续项目替换应以这 14 个一级目录为准。",
        ],
    )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    (OUTPUT_DIR / "company-level.md").write_text(company_markdown, encoding="utf-8")
    (OUTPUT_DIR / "department-level.md").write_text(department_markdown, encoding="utf-8")


if __name__ == "__main__":
    main()
